def run():
    import streamlit as st
    import pandas as pd
    import json
    from io import BytesIO
    from openpyxl import load_workbook
    from openpyxl.styles import PatternFill, Font
    from openai import OpenAI
    import pdfplumber
    from pdf2image import convert_from_bytes
    import pytesseract
    from PIL import Image
    import docx

    # ----------------- SETUP -----------------
    st.set_page_config(page_title="AI Resume Screener", page_icon="🤖", layout="centered")
    client = OpenAI(api_key=st.secrets["OPENAI_API_KEY"])

    # ----------------- TEXT EXTRACTION -----------------
    def extract_text_from_pdf(file):
        text = ""
        try:
            with pdfplumber.open(file) as pdf:
                for p in pdf.pages:
                    t = p.extract_text()
                    if t:
                        text += t + "\n"
        except Exception:
            pass
        if not text.strip():
            try:
                file.seek(0)
                imgs = convert_from_bytes(file.read())
                for img in imgs:
                    text += pytesseract.image_to_string(img)
            except Exception:
                text = ""
        return text.strip()

    def extract_text_from_docx(file):
        doc = docx.Document(file)
        return "\n".join(p.text for p in doc.paragraphs)

    # ----------------- LLM ANALYSIS -----------------
    def analyze_resume_with_llm(role, resume_text):
        prompt = f"""
You are an expert HR recruiter assistant.
Candidate is applying for {role}.
Below is their resume text.

1. Extract: Age, Education, Skills, Projects, Certifications.
2. Score each (0–10) for relevance to {role}.
3. Compute weighted average: Skills40 + Projects30 + Education20 + Certs10.
4. Return JSON only:
{{"weighted_average": float, "verdict": "PASS"/"FAIL", "reasoning": "string"}}
"""
        try:
            resp = client.responses.create(
                model="gpt-4o-mini",
                input=[
                    {"role": "system", "content": "Return valid JSON only."},
                    {"role": "user", "content": prompt + "\nResume:\n" + resume_text[:8000]},
                ],
                response_format={"type": "json_object"},
            )
            return resp.output_text
        except TypeError:
            chat = client.chat.completions.create(
                model="gpt-4o-mini",
                messages=[
                    {"role": "system", "content": "Return valid JSON only."},
                    {"role": "user", "content": prompt + "\nResume:\n" + resume_text[:8000]},
                ],
            )
            return chat.choices[0].message.content
        except Exception as e:
            return json.dumps({"weighted_average": 0, "verdict": "FAIL", "reasoning": f"Error: {e}"})

    # ----------------- APP LAYOUT -----------------
    st.title("🤖 AI Resume Screening System")

    if st.button("🔄 Reset Session"):
        st.session_state.clear()
        st.experimental_rerun()

    user_type = st.radio("Who are you?", ["Applicant", "HR Manager"])

    # ================== APPLICANT MODE ==================
    if user_type == "Applicant":
        st.subheader("👤 Applicant Mode")

        role = st.text_input("Enter the role you are applying for:")
        uploaded = st.file_uploader("Upload your resume (PDF or DOCX)", type=["pdf", "docx"])

        if uploaded and role:
            text = extract_text_from_pdf(uploaded) if uploaded.name.endswith(".pdf") else extract_text_from_docx(uploaded)
            st.text_area("📄 Resume Preview:", text[:1000])

            if text:
                with st.spinner("Analyzing your resume..."):
                    output = analyze_resume_with_llm(role, text)
                    try:
                        parsed = json.loads(output)
                        st.subheader("📊 Resume Screening Result")
                        st.json(parsed)
                        score = parsed.get("weighted_average", 0)
                        if score >= 6:
                            st.success("✅ You passed the screening!")
                        else:
                            st.error("❌ You did not pass the screening.")
                    except json.JSONDecodeError:
                        st.error("⚠️ Could not parse AI response.")
            else:
                st.warning("Could not read text from the uploaded file.")

    # ================== HR MANAGER MODE ==================
    elif user_type == "HR Manager":
        st.subheader("🧑‍💼 HR Manager Mode")
        role = st.text_input("Enter the role you are hiring for:")
        uploaded_files = st.file_uploader("Upload multiple resumes", type=["pdf", "docx"], accept_multiple_files=True)

        if uploaded_files and role:
            if "results" not in st.session_state or st.session_state.get("last_role") != role:
                total_files = len(uploaded_files)
                results = []
                failed_files = []

                st.info(f"Processing {total_files} resumes for the role of '{role}'...")
                progress = st.progress(0)
                status_text = st.empty()

                # 🔹 Pass 1: Initial Analysis
                for i, file in enumerate(uploaded_files):
                    with st.spinner(f"Analyzing {file.name} (Pass 1)..."):
                        text = extract_text_from_pdf(file) if file.name.endswith(".pdf") else extract_text_from_docx(file)
                        if text.strip():
                            llm_output = analyze_resume_with_llm(role, text)
                            try:
                                parsed = json.loads(llm_output)
                                parsed["filename"] = file.name
                                results.append(parsed)
                            except json.JSONDecodeError:
                                failed_files.append(file)
                        else:
                            failed_files.append(file)

                    progress.progress((i + 1) / total_files)
                    status_text.text(f"✅ Processed {i + 1}/{total_files} resumes (Pass 1).")

                # 🔁 Pass 2: Retry Failed Resumes Automatically
                if failed_files:
                    st.warning(f"⚠️ Retrying {len(failed_files)} failed resumes automatically...")
                    for j, file in enumerate(failed_files):
                        with st.spinner(f"Re-analyzing {file.name} (Retry {j+1}/{len(failed_files)})..."):
                            text = extract_text_from_pdf(file) if file.name.endswith(".pdf") else extract_text_from_docx(file)
                            if text.strip():
                                llm_output = analyze_resume_with_llm(role, text)
                                try:
                                    parsed = json.loads(llm_output)
                                    parsed["filename"] = file.name
                                    results.append(parsed)
                                    st.success(f"✅ {file.name} parsed successfully on retry.")
                                except json.JSONDecodeError:
                                    st.error(f"❌ {file.name} could not be parsed after retry.")
                                    results.append({
                                        "filename": file.name,
                                        "weighted_average": 0,
                                        "verdict": "FAIL",
                                        "reasoning": "Resume could not be parsed after two attempts."
                                    })
                            else:
                                st.error(f"❌ {file.name} unreadable (even after retry).")
                                results.append({
                                    "filename": file.name,
                                    "weighted_average": 0,
                                    "verdict": "FAIL",
                                    "reasoning": "Unreadable or empty resume text (after retry)."
                                })

                # 🧾 Save All Results in Session
                st.session_state["results"] = results
                st.session_state["last_role"] = role

            # Retrieve results from session
            results = st.session_state.get("results", [])

            # ✅ Final Excel Export
            if results:
                df = pd.DataFrame(results)
                df = df.sort_values(by="weighted_average", ascending=False)

                st.success("✅ Screening complete (including retries)!")
                st.dataframe(df[["filename", "weighted_average", "verdict", "reasoning"]])

                # 🧠 Summary info for HR
                st.info(f"Total resumes processed: {len(uploaded_files)} | Final results: {len(results)}")

                # 🏅 Leaderboard (Top 5)
                st.subheader("🏅 Top 5 Candidates")
                top5 = df.head(5)
